import io

import docx
import pytest
from fastapi import UploadFile
from PIL import Image, ImageDraw, ImageFont

from app.core.exceptions import BadRequestError, PayloadTooLargeError
from app.modules.ai import documents
from app.modules.ai.documents import MAX_TEXT_CHARS, extract_text, read_upload_text


def text_image(lines: list[str], fmt: str = "PNG") -> bytes:
    """A white image with large black text, like a business card photo."""
    image = Image.new("RGB", (900, 120 + 80 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=40)
    for i, line in enumerate(lines):
        draw.text((40, 40 + i * 80), line, fill="black", font=font)
    output = io.BytesIO()
    image.save(output, format=fmt)
    return output.getvalue()


def text_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a real text layer."""
    content = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = output.tell()
    output.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        output.write(b"%010d 00000 n \n" % offset)
    output.write(b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1))
    output.write(b"startxref\n%d\n%%%%EOF\n" % xref)
    return output.getvalue()


def test_image_is_read_with_ocr() -> None:
    text = extract_text("card.png", "image/png", text_image(["Dana Lee", "dana@acme.io"]))

    assert "Dana Lee" in text
    assert "dana@acme.io" in text


def test_pdf_text_layer() -> None:
    assert extract_text("intro.pdf", "application/pdf", text_pdf("Jane Doe CTO")) == "Jane Doe CTO"


def test_scanned_pdf_falls_back_to_ocr() -> None:
    scan = io.BytesIO()
    Image.open(io.BytesIO(text_image(["Scanned Person"]))).save(scan, format="PDF")

    assert "Scanned Person" in extract_text("scan.pdf", "application/pdf", scan.getvalue())


def test_docx_paragraphs_and_tables() -> None:
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Email"
    table.rows[0].cells[1].text = "jane@acme.com"
    output = io.BytesIO()
    document.save(output)

    text = extract_text("cv.docx", "application/octet-stream", output.getvalue())

    assert "Jane Doe" in text
    assert "Email | jane@acme.com" in text


def test_email_headers_and_body() -> None:
    raw = b"From: Jane Doe <jane@acme.com>\nSubject: Intro\n\nHi there,\nJane, CTO at Acme\n"

    text = extract_text("intro.eml", "message/rfc822", raw)

    assert "From: Jane Doe <jane@acme.com>" in text
    assert "Subject: Intro" in text
    assert "CTO at Acme" in text


def test_plain_text_and_vcard() -> None:
    assert extract_text("notes.txt", "text/plain", b"Jane Doe") == "Jane Doe"
    assert "FN:Jane" in extract_text("jane.vcf", "", b"BEGIN:VCARD\nFN:Jane\nEND:VCARD")


def test_long_text_is_truncated() -> None:
    text = extract_text("long.txt", "text/plain", b"a" * (MAX_TEXT_CHARS + 100))
    assert len(text) == MAX_TEXT_CHARS


@pytest.mark.parametrize(
    ("filename", "content_type", "data", "message"),
    [
        ("archive.zip", "application/zip", b"PK", "Unsupported file"),
        ("blank.txt", "text/plain", b"   \n", "No text was found"),
        ("broken.png", "image/png", b"not an image", "image could not be read"),
        ("broken.pdf", "application/pdf", b"%PDF-garbage", "PDF could not be read"),
        ("broken.docx", "", b"not a zip", "Word document could not be read"),
    ],
)
def test_bad_files_are_rejected(
    filename: str, content_type: str, data: bytes, message: str
) -> None:
    with pytest.raises(BadRequestError, match=message):
        extract_text(filename, content_type, data)


async def test_read_upload_text_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents, "MAX_FILE_BYTES", 10)

    with pytest.raises(PayloadTooLargeError):
        await read_upload_text(UploadFile(io.BytesIO(b"x" * 11), filename="a.txt"))
    with pytest.raises(BadRequestError, match="empty"):
        await read_upload_text(UploadFile(io.BytesIO(b""), filename="a.txt"))

    upload = UploadFile(io.BytesIO(b"Jane"), filename="a.txt")
    assert await read_upload_text(upload) == "Jane"
